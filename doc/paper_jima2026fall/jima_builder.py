#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""日本経営工学会 予稿テンプレート(jima_abst_template.docx)へ原稿を流し込む docx ビルダー。

テンプレート実測仕様:
  - A4 縦・余白 20mm 四方・docGrid linePitch=291twips（1頁目 45行 / 2頁目以降 50行）
  - 題目部は 1 段セクション、本文以降は 2 段セクション（段間 398twips, 段幅 4620twips）
  - 段落スタイルはテンプレ同梱の Jima_* を使う（見出しは numId=1 で自動採番）

コンテンツはブロックのリスト（content_ja.py 参照）:
  ('title', str) ('affil_star', str) ('affil', str)
  ('h1', str) ('h2', str) ('p', str)
  ('eq', 'EQ1'|'EQ2', '(1)')
  ('fig', png, caption, 'col'|'full', scale)
  ('table', header, rows, widths, caption)
  ('refs_heading', str) ('ref', str)
インライン記法は APIEMS 版 builder と共通: **bold** / *italic* / $math$
"""
import copy
import os
import sys

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Emu, Twips

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.append(os.path.normpath(os.path.join(HERE, '..', 'paper_apiems2026')))
import builder as ab  # noqa: E402  （rich_runs / OMML 数式を再利用）

NS_W = ab.NS_W
NS_M = ab.NS_M

# 段組寸法（テンプレ実測, twips）
PAGE_W, MAR = 11906, 1134
COL_GAP = 398
TEXT_W = PAGE_W - MAR * 2                    # 9638
COL_W = (TEXT_W - COL_GAP) // 2              # 4620

SECT_PGDIMS = ('<w:pgSz w:w="11906" w:h="16838" w:code="9"/>'
               '<w:pgMar w:top="1134" w:right="1134" w:bottom="1134"'
               ' w:left="1134" w:header="851" w:footer="992" w:gutter="0"/>')

EA = 'ＭＳ Ｐ明朝'
REF_SZ = 18   # 参考文献は 9pt（half-points）


def sect_xml(cols):
    if cols == 2:
        c = f'<w:cols w:num="2" w:space="{COL_GAP}"/>'
        grid = '<w:docGrid w:type="linesAndChars" w:linePitch="291"/>'
    else:
        c = '<w:cols w:space="425"/>'
        grid = '<w:docGrid w:type="lines" w:linePitch="291"/>'
    return (f'<w:sectPr {NS_W}><w:type w:val="continuous"/>'
            f'{SECT_PGDIMS}{c}{grid}</w:sectPr>')


def p_xml(inner, ppr=''):
    return f'<w:p {NS_W} {NS_M}><w:pPr>{ppr}</w:pPr>{inner}</w:p>'


def para_style(text, style, extra_ppr='', bold=False, size=None):
    runs = '<w:r><w:br/></w:r>'.join(
        ab.rich_runs(t, EA, base_bold=bold) for t in text.split('\n'))
    if size:
        runs = runs.replace('</w:rPr>',
                            f'<w:sz w:val="{size}"/><w:szCs w:val="{size}"/></w:rPr>')
    return p_xml(runs, f'<w:pStyle w:val="{style}"/>{extra_ppr}')


def para_body(text):
    # テンプレ本文は全角1字下げを直書きする
    if not text.startswith('　'):
        text = '　' + text
    return para_style(text, 'ima010')


def para_h1(text, first=False):
    return para_style(text, 'ima004' if first else 'Jima005')


def para_h2(text):
    return para_style(text, 'Jima005',
                      '<w:numPr><w:ilvl w:val="1"/><w:numId w:val="1"/></w:numPr>')


def para_caption(text):
    return para_style(text, 'ima200', '<w:keepLines/>')


def para_sect_break(cols):
    return (f'<w:p {NS_W}><w:pPr>'
            f'<w:spacing w:line="14" w:lineRule="exact"/>'
            f'{sect_xml(cols)}'
            f'<w:rPr><w:sz w:val="2"/></w:rPr></w:pPr></w:p>')


def para_equation(omml, number, width=None):
    w = width or COL_W
    tabs = (f'<w:tabs><w:tab w:val="center" w:pos="{w // 2}"/>'
            f'<w:tab w:val="right" w:pos="{w}"/></w:tabs>')
    inner = ('<w:r><w:tab/></w:r>' + omml + '<w:r><w:tab/></w:r>' +
             ab._run(number, east_asia=EA))
    return p_xml(inner, '<w:spacing w:before="60" w:after="60"/>' + tabs)


def table_xml(header, rows, widths, total_w=COL_W, fs=15):
    tw = [int(total_w * w / sum(widths)) for w in widths]

    def cell(text, w, bold=False, top=False, bottom=False, center=True):
        borders = '<w:tcBorders>'
        if top:
            borders += '<w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        if bottom:
            borders += '<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        borders += '</w:tcBorders>'
        runs = ab.rich_runs(text, EA, base_bold=bold).replace(
            '</w:rPr>', f'<w:sz w:val="{fs}"/><w:szCs w:val="{fs}"/></w:rPr>')
        jc = '<w:jc w:val="center"/>' if center else ''
        return (f'<w:tc><w:tcPr><w:tcW w:w="{w}" w:type="dxa"/>{borders}</w:tcPr>'
                f'<w:p><w:pPr><w:spacing w:line="180" w:lineRule="exact"/>'
                f'{jc}</w:pPr>{runs}</w:p></w:tc>')

    xml = (f'<w:tbl {NS_W}><w:tblPr><w:tblW w:w="{total_w}" w:type="dxa"/>'
           '<w:jc w:val="center"/><w:tblLayout w:type="fixed"/>'
           '<w:tblCellMar><w:left w:w="28" w:type="dxa"/>'
           '<w:right w:w="28" w:type="dxa"/></w:tblCellMar>'
           '</w:tblPr><w:tblGrid>' +
           ''.join(f'<w:gridCol w:w="{w}"/>' for w in tw) + '</w:tblGrid>')
    xml += '<w:tr>' + ''.join(
        cell(h, tw[j], bold=False, top=True, bottom=True)
        for j, h in enumerate(header)) + '</w:tr>'
    for i, row in enumerate(rows):
        last = i == len(rows) - 1
        xml += '<w:tr>' + ''.join(
            cell(c, tw[j], bottom=last) for j, c in enumerate(row)) + '</w:tr>'
    xml += '</w:tbl>'
    return xml


def build(template_path, out_path, blocks, fig_dir=''):
    doc = Document(template_path)
    body = doc.element.body

    # テンプレ先頭（1 段・題目部）の sectPr を確保
    title_sect = None
    for p in body.iter(qn('w:p')):
        pPr = p.find(qn('w:pPr'))
        if pPr is not None and pPr.find(qn('w:sectPr')) is not None:
            title_sect = copy.deepcopy(pPr.find(qn('w:sectPr')))
            break
    if title_sect is None:
        raise RuntimeError('template first sectPr not found')

    for el in list(body):
        if el.tag != qn('w:sectPr'):
            body.remove(el)
    final_sect = body.find(qn('w:sectPr'))
    for el in list(final_sect):
        final_sect.remove(el)
    for el in list(parse_xml(sect_xml(2))):
        final_sect.append(copy.deepcopy(el))

    def add(xml_str):
        el = parse_xml(xml_str)
        body.insert(len(body) - 1, el)
        return el

    def add_picture(path, box_w, scale):
        from PIL import Image
        img = Image.open(os.path.join(fig_dir, path))
        w_px, h_px = img.size
        w_emu = int(Twips(int(box_w * scale)))
        h_emu = int(w_emu * h_px / w_px)
        pic_p = doc.add_paragraph()
        ppr = pic_p._p.get_or_add_pPr()
        ppr.append(parse_xml(f'<w:keepNext {NS_W}/>'))
        ppr.append(parse_xml(f'<w:spacing {NS_W} w:line="240" w:lineRule="auto"'
                             ' w:before="60" w:after="0"/>'))
        ppr.append(parse_xml(f'<w:jc {NS_W} w:val="center"/>'))
        pic_p.add_run().add_picture(os.path.join(fig_dir, path),
                                    width=Emu(w_emu), height=Emu(h_emu))

    first_h1 = True
    for blk in blocks:
        kind = blk[0]
        if kind == 'title':
            add(para_style(blk[1], 'ima001'))
        elif kind == 'affil_star':
            add(para_style(blk[1], 'ima002'))
        elif kind == 'affil':
            # 規定（本文は第 6 行目以降）に合わせ、所属行のあとの空きを 1 行に抑える
            add(para_style(blk[1], 'ima003',
                           '<w:spacing w:afterLines="50" w:after="50"/>'))
            # 題目部（1 段セクション）をここで閉じる
            p = add(f'<w:p {NS_W}><w:pPr>'
                    f'<w:spacing w:line="14" w:lineRule="exact"/>'
                    f'<w:rPr><w:sz w:val="2"/></w:rPr></w:pPr></w:p>')
            p.find(qn('w:pPr')).append(title_sect)
        elif kind == 'h1':
            add(para_h1(blk[1], first=first_h1))
            first_h1 = False
        elif kind == 'h2':
            add(para_h2(blk[1]))
        elif kind == 'p':
            add(para_body(blk[1]))
        elif kind == 'eq':
            omml = ab.omml_eq1() if blk[1] == 'EQ1' else ab.omml_eq2()
            add(para_equation(omml, blk[2]))
        elif kind == 'fig':
            path, caption, mode = blk[1], blk[2], blk[3]
            scale = blk[4] if len(blk) > 4 else 1.0
            if mode == 'full':
                add(para_sect_break(2))
                add_picture(path, TEXT_W, scale)
                add(para_caption(caption))
                add(para_sect_break(1))
            else:
                add_picture(path, COL_W, scale)
                add(para_caption(caption))
        elif kind == 'table':
            header, rows, widths, caption = blk[1], blk[2], blk[3], blk[4]
            add(para_caption(caption))
            add(table_xml(header, rows, widths))
            add(f'<w:p {NS_W}><w:pPr>'
                f'<w:spacing w:line="120" w:lineRule="exact"/>'
                f'<w:rPr><w:sz w:val="12"/></w:rPr></w:pPr></w:p>')
        elif kind == 'refs_heading':
            add(para_style(blk[1], 'ima006',
                           '<w:numPr><w:ilvl w:val="0"/>'
                           '<w:numId w:val="0"/></w:numPr>'
                           '<w:spacing w:before="60"/>'))
        elif kind == 'ref':
            add(para_style(blk[1], 'ima100References',
                           '<w:snapToGrid w:val="0"/>'
                           '<w:spacing w:line="200" w:lineRule="exact"/>'
                           '<w:ind w:left="420" w:hanging="420"/>'
                           '<w:jc w:val="left"/>',
                           size=REF_SZ))
        else:
            raise ValueError(f'unknown block: {kind}')

    doc.save(out_path)
    print(' ->', out_path)
